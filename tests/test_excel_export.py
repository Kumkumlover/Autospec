"""Unit tests for Excel Export & Formatting (boq_engine.py)."""

import openpyxl
import pytest

from boq_engine import BoqEngine
from cad_parser import parse_dxf_file
from sample_dxf_generator import generate_sample_2bhk_dxf
from spec_parser import parse_client_brief


@pytest.fixture(scope="session")
def generated_boq_excel(tmp_path_factory):
    dxf_path = str(tmp_path_factory.mktemp("dxf") / "sample.dxf")
    excel_path = str(tmp_path_factory.mktemp("excel") / "output_boq.xlsx")

    generate_sample_2bhk_dxf(dxf_path)
    takeoff = parse_dxf_file(dxf_path)
    spec = parse_client_brief("Luxury 2BHK with Philips 12W 3000K spots and Atomberg BLDC fans.")

    engine = BoqEngine()
    boq = engine.generate_boq(takeoff, spec)
    engine.export_to_excel(boq, excel_path)
    return excel_path


def test_excel_export_sheets_exist(generated_boq_excel):
    wb = openpyxl.load_workbook(generated_boq_excel, data_only=False)
    assert "Priced BOQ (IS 1200)" in wb.sheetnames
    assert "IS 1200 Deduction Audit" in wb.sheetnames


def test_excel_export_formulas_intact(generated_boq_excel):
    wb = openpyxl.load_workbook(generated_boq_excel, data_only=False)
    ws = wb["Priced BOQ (IS 1200)"]

    # Check for formula in Total Cost column (Col 7 / G)
    has_product_formula = False
    has_sum_formula = False

    for r in range(5, ws.max_row + 1):
        cell_val = str(ws.cell(row=r, column=7).value or "")
        if cell_val.startswith("=D") and "*" in cell_val:
            has_product_formula = True
        if cell_val.startswith("=SUM("):
            has_sum_formula = True

    assert has_product_formula, "Line items should have dynamic product formula =D*F"
    assert has_sum_formula, "Grand total row should have =SUM() formula"


def test_excel_export_hyperlinks(generated_boq_excel):
    wb = openpyxl.load_workbook(generated_boq_excel, data_only=False)
    ws = wb["Priced BOQ (IS 1200)"]

    # Col 8 is Buy URL link
    links_found = 0
    for r in range(5, ws.max_row):
        cell = ws.cell(row=r, column=8)
        if cell.hyperlink and cell.hyperlink.target:
            assert cell.hyperlink.target.startswith("http")
            links_found += 1

    assert links_found > 0, "Procurement items must contain active hyperlinks"
